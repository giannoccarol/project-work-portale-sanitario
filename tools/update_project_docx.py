"""Aggiorna il DOCX finale mantenendo stile, immagini e diagrammi.

Per una nuova revisione aggiungere una coppia testo-prima/testo-dopo a
``REPLACEMENTS`` e lanciare:

    python tools/update_project_docx.py percorso/al/documento.docx

Il controllo di cardinalità impedisce di sostituire un paragrafo ambiguo.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


# Inserire qui solo sostituzioni riferite alla versione corrente del documento.
REPLACEMENTS: dict[str, str] = {}


def iter_paragraphs(document: Document):
    seen: set[int] = set()
    for element in document.element.body.iter(qn("w:p")):
        if id(element) in seen:
            continue
        seen.add(id(element))
        yield Paragraph(element, document)
    for section in document.sections:
        for paragraph in (*section.header.paragraphs, *section.footer.paragraphs):
            if id(paragraph._p) in seen:
                continue
            seen.add(id(paragraph._p))
            yield paragraph


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Uso: update_project_docx.py <documento.docx>")

    target = Path(sys.argv[1]).resolve()
    document = Document(target)
    paragraphs = list(iter_paragraphs(document))
    modifications = 0

    for old, new in REPLACEMENTS.items():
        old_matches = [paragraph for paragraph in paragraphs if paragraph.text == old]
        new_matches = [paragraph for paragraph in paragraphs if paragraph.text == new]
        if len(old_matches) == 1 and not new_matches:
            paragraph = old_matches[0]
            if paragraph.runs:
                paragraph.runs[0].text = new
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(new)
            modifications += 1
        elif not old_matches and len(new_matches) == 1:
            continue
        else:
            raise RuntimeError(
                f"Sostituzione non conforme: testo prima={len(old_matches)}, "
                f"testo dopo={len(new_matches)} ({old[:80]!r})"
            )

    document.core_properties.title = "Puglia Salute - Portale sanitario full-stack"
    document.core_properties.subject = (
        "Project Work PW16: prenotazioni e referti nella rete sanitaria pugliese"
    )
    document.core_properties.keywords = (
        "Puglia Salute, Angular 21, PrimeNG 21, Express, TypeORM, "
        "prenotazioni, referti"
    )

    temporary = target.with_suffix(".docx.tmp")
    document.save(temporary)
    os.replace(temporary, target)
    print(f"Aggiornato {target.name}: {modifications} paragrafi sostituiti")


if __name__ == "__main__":
    main()
